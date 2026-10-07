"""Gera BD.xlsx: base fictícia e original, no layout que o Power Query do relatório espera."""
import numpy as np, pandas as pd, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font
rng = np.random.default_rng(2026)

lojas = pd.DataFrame({"codigo_loja": ["CL001", "CL002", "CL003", "CL004"],
                      "nome_loja": ["Matriz", "Filial MG", "Filial SP", "Filial ES"],
                      "cidade_loja": ["Rio de Janeiro", "Belo Horizonte", "São Paulo", "Vitória"]})

prods = [  # descrição, categoria, custo, preço
 ("Notebook 15", "Eletrônicos", 2150.0, 3399.9), ("Smartwatch", "Eletrônicos", 520.0, 899.9),
 ("Fone Bluetooth", "Eletrônicos", 95.0, 219.9), ("Caixa de Som Portátil", "Eletrônicos", 180.0, 379.9),
 ("Monitor 27", "Eletrônicos", 740.0, 1249.9),
 ("Geladeira Frost Free", "Eletrodomésticos", 1890.0, 2999.9), ("Fogão 5 Bocas", "Eletrodomésticos", 780.0, 1349.9),
 ("Micro-ondas 30L", "Eletrodomésticos", 390.0, 649.9), ("Aspirador Robô", "Eletrodomésticos", 610.0, 1099.9),
 ("Ar-condicionado 12000", "Eletrodomésticos", 1420.0, 2299.9),
 ("Air Fryer 5L", "Cozinha", 240.0, 449.9), ("Cafeteira Expresso", "Cozinha", 330.0, 599.9),
 ("Liquidificador Turbo", "Cozinha", 110.0, 229.9), ("Jogo de Panelas", "Cozinha", 160.0, 329.9),
 ("Batedeira Planetária", "Cozinha", 290.0, 529.9),
 ("Colchão Queen", "Casa", 980.0, 1699.9), ("Poltrona Reclinável", "Casa", 690.0, 1199.9),
 ("Ventilador de Coluna", "Casa", 130.0, 259.9), ("Luminária de Mesa", "Casa", 45.0, 109.9),
 ("Jogo de Cama Casal", "Casa", 85.0, 189.9)]
dprod = pd.DataFrame([(20001 + i, d, c) for i, (d, c, _, _) in enumerate(prods)],
                     columns=["codigo_produto", "descricao_produto", "Categoria"])
custo = {20001 + i: p[2] for i, p in enumerate(prods)}; preco = {20001 + i: p[3] for i, p in enumerate(prods)}
peso_prod = np.array([3, 6, 9, 7, 4, 3, 4, 7, 4, 3, 8, 6, 8, 6, 5, 3, 3, 7, 8, 8], float); peso_prod /= peso_prod.sum()

nomes = [("Aline Barreto", "F"), ("Caio Nogueira", "M"), ("Débora Pacheco", "F"),
         ("Everton Salles", "M"), ("Fernanda Quintela", "F"), ("Gustavo Meireles", "M"),
         ("Helena Fragoso", "F"), ("Igor Bastos", "M"), ("Jéssica Andrade", "F"),
         ("Leandro Viana", "M"), ("Marina Tavares", "F"), ("Otávio Cordeiro", "M")]
perf_op = ["Acima do Esperado", "Dentro do Esperado", "Abaixo do Esperado"]
func = []
for i, (n, s) in enumerate(nomes):
    perf = perf_op[[0, 1, 2, 1, 0, 1, 0, 2, 1, 1, 0, 2][i]]
    func.append({
        "nome_funcionario": n, "matricula_funcionario": 880001 + i, "PerfScoreID": 4 - perf_op.index(perf),
        "salario_mensal": int(rng.integers(30, 62) * 100), "Cargo": "Vendedor",
        "dt_nascimento": dt.date(int(rng.integers(1975, 2000)), int(rng.integers(1, 13)), int(rng.integers(1, 28))),
        "Sexo": s, "estado_civil": rng.choice(["Solteiro", "Casado"]),
        "dt_contratacao": dt.date(int(rng.integers(2014, 2019)), int(rng.integers(1, 13)), int(rng.integers(1, 28))),
        "dt_demissao": None, "motivo_saida": None, "status_funcionario": "Ativo",
        "Departamento": f"LOJA 0{i // 3 + 1}",
        "fonte_recrutamento": rng.choice(["Site de Vagas", "Site da Empresa", "Indicação Funcionários", "Feira de Contratação"]),
        "performance": perf, "pesquisa_engajamento": round(float(rng.uniform(1.8, 5)), 2),
        "indice_satisfacao": float([4, 3, 2, 5, 4, 3, 5, 1, 4, 3, 5, 2][i]),
        "ultima_atualizacao_performance": dt.date(2020, 12, 31)})
dfunc = pd.DataFrame(func)

# fVendas: 3000 vendas em 2019-2020, com sazonalidade, crescimento em 2020 e vendedores de desempenho diferente
dias = pd.date_range("2019-01-01", "2020-12-31")
saz = {1: .9, 2: .85, 3: .95, 4: 1, 5: 1.15, 6: 1, 7: .95, 8: 1.05, 9: 1, 10: 1.05, 11: 1.4, 12: 1.5}
w = np.array([saz[d.month] * (1.18 if d.year == 2020 else 1) * (1.25 if d.dayofweek >= 4 else 1) for d in dias]); w /= w.sum()
forca = np.array([1.9, 1.0, .7, 1.2, 1.75, .8, 1.3, .6, 1.0, .9, 1.6, .65]); 
N = 3000
d_venda = rng.choice(dias, N, p=w)
vend = rng.choice(12, N, p=forca / forca.sum())
cod = rng.choice(dprod.codigo_produto, N, p=peso_prod)
seg = rng.integers(8 * 3600 + 60, 20 * 3600, N)
atraso = rng.choice([0, 1, 2, 3, 4, 5, 6, 8], N, p=[.3, .25, .15, .1, .08, .06, .04, .02])
seg_ent = rng.integers(8 * 3600 + 60, 19 * 3600, N)
rows = []
for i in range(N):
    dv = pd.Timestamp(d_venda[i]) + pd.Timedelta(seconds=int(seg[i]))
    de = pd.Timestamp(d_venda[i]) + pd.Timedelta(days=int(atraso[i]), seconds=int(seg_ent[i]))
    if de.year > 2020: de = pd.Timestamp(d_venda[i]) + pd.Timedelta(seconds=int(seg_ent[i]))  # entrega fica dentro de 2020
    if de <= dv: de = dv + pd.Timedelta(minutes=int(rng.integers(20, 180)))
    f = dfunc.iloc[vend[i]]
    rows.append([int(f.matricula_funcionario), f.nome_funcionario, "Vendedor", lojas.codigo_loja[vend[i] // 3], int(cod[i]),
                 dprod.set_index("codigo_produto").Categoria[cod[i]], custo[cod[i]], preco[cod[i]],
                 int(rng.choice([1, 2, 3, 4], p=[.34, .28, .22, .16])), dv.to_pydatetime(), de.to_pydatetime()])
fv = pd.DataFrame(rows, columns=["matricula_funcionario", "Nome Funcionario", "Cargo", "codigo_loja", "codigo_produto", "Categoria",
                                 "preco_custo", "valor_unitario", "quantidade", "dt_venda", "dt_entrega"]).sort_values("dt_venda").reset_index(drop=True)

# Metas: 2019 fixa por loja; 2020 = realizado do mesmo mês de 2019 + 12%
fv["fat"] = fv.valor_unitario * fv.quantidade
real19 = fv[fv.dt_venda.dt.year == 2019].groupby(["codigo_loja", fv.dt_venda.dt.month]).fat.sum()
meses = [f"01/{m:02d}/{a}" for a in (2019, 2020) for m in range(1, 13)]
metas = []
for c in lojas.codigo_loja:
    m19 = round(real19[c].sum() * 0.97 / 12, -2)
    v = [m19] * 12 + [round(real19[c][m] * 1.12, 2) for m in range(1, 13)]
    metas.append([c] + v + [round(sum(v[:12]), 2), round(sum(v[12:]), 2)])
tot = ["Total"] + [round(sum(r[j] for r in metas), 2) for j in range(1, 27)]

wb = Workbook(); wb.remove(wb.active)
def folha(nome, df):
    ws = wb.create_sheet(nome); ws.append(list(df.columns))
    for r in df.itertuples(index=False): ws.append(list(r))
    return ws
ws = folha("fVendas", fv.drop(columns="fat"))
for r in ws.iter_rows(min_row=2, min_col=10, max_col=11):
    for c in r: c.number_format = "dd/mm/yyyy hh:mm:ss"
ws = folha("dFuncionario", dfunc)
for col in (6, 9, 18):
    for r in ws.iter_rows(min_row=2, min_col=col, max_col=col):
        r[0].number_format = "dd/mm/yyyy"
folha("dLoja", lojas); folha("dProdutos", dprod)
ws = wb.create_sheet("Metas"); ws.append(["Meta de Faturamento"])
ws.append(["Orcamento"] + meses + ["Meta 2019", "Meta 2020"])
for r in metas + [tot]: ws.append(r)
for ws in wb:
    for row in ws.iter_rows():
        for c in row: c.font = Font(name="Arial", size=10, bold=(c.row == 1 or (ws.title == "Metas" and c.row == 2)))
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(28, max(len(str(c.value or "")) for c in col[:50]) + 2)
    ws.freeze_panes = "A3" if ws.title == "Metas" else "A2"
wb.save("/mnt/user-data/outputs/BD.xlsx")

# conferência: simula as medidas principais
print("faturamento", round(fv.fat.sum(), 2), "| pedidos", len(fv), "| ticket", round(fv.fat.mean(), 2))
print("lucro", round(((fv.valor_unitario - fv.preco_custo) * fv.quantidade).sum(), 2))
print(fv.groupby(fv.dt_venda.dt.year).fat.sum().round(0).to_dict(), "| meta", round(tot[25]), round(tot[26]))
print("por vendedor:", fv.groupby("Nome Funcionario").fat.sum().round(0).sort_values().to_dict())
print("ES", round(fv[fv.codigo_loja == "CL004"].fat.sum()), "| qtd>=3:", (fv.quantidade >= 3).sum(), "| meia-noite:", ((fv.dt_venda.dt.hour == 0) | (fv.dt_entrega.dt.hour == 0)).sum(), "| entrega<venda:", (fv.dt_entrega < fv.dt_venda).sum())
